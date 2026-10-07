#!/usr/bin/env python3
"""Verify the real reporting CLI uses all inputs and identifies historical slides."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
from pptx import Presentation

script = Path(__file__).resolve().parents[2] / "scripts/build_presentation.py"
with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    (root / "molecules").mkdir()
    (root / "outs").mkdir()
    (root / "molecules/S1.molecules.metrics.json").write_text(json.dumps({
        "eligible_records": 987, "molecules_counted": 37, "representatives_tagged": 38,
        "fractional_collisions": True,
    }))
    (root / "outs/cell_calling_metrics.json").write_text(json.dumps({
        "filtered_cells": 4, "filtered_molecules": 29,
    }))
    illumina = root / "illumina_metrics.json"
    illumina.write_text(json.dumps({"overlap_enabled": True, "overlap_short_sins": 3,
        "overlap_medium_mins": 1, "overlap_long_lins": 2,
        "insert_length_histogram": {"120": 3, "180": 1}}))
    for mode, flags in [
        ("discover", ["--report-dir", str(root)]),
        ("explicit", ["--illumina-metrics", str(illumina), "--molecules-metrics", str(root / "molecules/S1.molecules.metrics.json"),
                      "--cell-metrics", str(root / "outs/cell_calling_metrics.json")]),
    ]:
        output = root / f"{mode}.pptx"
        tsv = root / f"{mode}.tsv"
        subprocess.run([sys.executable, str(script), "--out", str(output), "--export-insert-sizes", str(tsv), *flags], check=True)
        deck = Presentation(output)
        text = "\n".join(
            shape.text if shape.has_text_frame else
            "\n".join(cell.text for row in shape.table.rows for cell in row.cells) if shape.has_table else ""
            for slide in deck.slides for shape in slide.shapes
        )
        for expected in ["Current run metrics", "Counted molecules", "37", "Called cells/barcodes", "4", "29", "Collision policy: fractional"]:
            assert expected in text, expected
        assert text.count("Historical mapping example") == 8
        assert tsv.read_text() == "insert_length\tcount\n120\t3\n180\t1\n"
